#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""在数学建模论文终稿后追加完整程序代码附录。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "运行结果_20260819"
INPUT_DOCX = RESULTS / "数学建模论文终稿.docx"
OUTPUT_DOCX = RESULTS / "数学建模论文终稿_含完整附录.docx"

EXTERNAL_SOURCE_DIR = ROOT / "external_plot_sources"


MODEL_SOURCES = [
    {
        "number": "B.1",
        "title": "问题一：库存平衡班次优化求解程序",
        "file": ROOT / "solve_question1.py",
        "purpose": "读取逐小时到货量，构建班次覆盖、处理能力与库存守恒约束，执行词典序混合整数规划并输出每日求解结果。",
    },
    {
        "number": "B.2",
        "title": "问题二：时限与低效约束排班求解程序",
        "file": ROOT / "solve_question2.py",
        "purpose": "在问题一基础上区分早期货物与普通货物，加入16:00截止、低效小时和共享产能约束，求解每日最低人数。",
    },
    {
        "number": "B.3",
        "title": "问题三：连续工日约束人员配置求解程序",
        "file": ROOT / "solve_question3.py",
        "purpose": "构建累计工日—连续工日二维状态流模型，计算招聘人数下界，求解聚合整数流并分解为具名员工月历。",
    },
    {
        "number": "B.4",
        "title": "问题三融合优化程序",
        "file": ROOT / "solve_question3_hybrid.py",
        "purpose": "联立日内排班与月度人员状态，对问题三方案进行融合求解和可执行性处理。",
    },
    {
        "number": "B.5",
        "title": "问题三结果后处理与优化程序",
        "file": ROOT / "optimize_question3_results.py",
        "purpose": "对月度排班解进行低效小时分配、具名员工路径整理与结果工作簿增强。",
    },
    {
        "number": "B.6",
        "title": "问题三独立核验程序",
        "file": ROOT / "verify_question3.py",
        "purpose": "逐项核对每人工日数、连续工日上限、每日覆盖、具名分解与问题二日内可行性。",
    },
    {
        "number": "B.7",
        "title": "三问模型有效性与鲁棒性检验程序",
        "file": ROOT / "validate_all_models.py",
        "purpose": "完成约束残差、整数性、理论下界、独立重求解、随机扰动和系统性压力测试。",
    },
]


FIGURE_SOURCES = [
    {
        "number": "C.1",
        "title": "附录图表批量绘制程序",
        "file": ROOT / "figures" / "gen_appendix_figures.py",
        "purpose": "批量生成三问人数对比、逐小时平均库存、制度性冗余和连续工日分布图。",
    },
    {
        "number": "C.2",
        "title": "平均库存与连续工日分布绘图程序",
        "file": EXTERNAL_SOURCE_DIR / "plot_inventory_and_streak.py",
        "purpose": "从求解工作簿读取真实数据，生成两种日排班模型的逐小时平均库存图和员工最大连续工作天数分布图。",
    },
    {
        "number": "C.3",
        "title": "三问每日人员规模绘图程序",
        "file": EXTERNAL_SOURCE_DIR / "redraw_staffing_figures.py",
        "purpose": "生成问题二最低需求与问题三实际出勤对比图，以及三问每日人员规模对比图。",
    },
    {
        "number": "C.4",
        "title": "三问能力利用率与人员冗余对比绘图程序",
        "file": EXTERNAL_SOURCE_DIR / "bule.py",
        "purpose": "绘制三个问题的处理能力利用率与月累计人日对比图。",
    },
    {
        "number": "C.5",
        "title": "处理效率敏感性绘图程序",
        "file": EXTERNAL_SOURCE_DIR / "last.py",
        "purpose": "展示正常处理效率上下浮动10%时三问人员规模的变化。",
    },
    {
        "number": "C.6",
        "title": "30天逐小时到货量热力图绘制程序",
        "file": EXTERNAL_SOURCE_DIR / "draw_fig2_heatmap_from_reference(1).py",
        "purpose": "生成30天×24小时到货量热力图，标注高、低负荷日及颜色尺度。",
    },
    {
        "number": "C.7",
        "title": "问题二低效小时与双截止期库存流示意图程序",
        "file": EXTERNAL_SOURCE_DIR / "draw_fig4_from_reference(1).py",
        "purpose": "绘制早期货物、普通货物、低效人数与共同产能池之间的关系示意图。",
    },
]


def read_text_safely(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", b"", 0, 1, f"无法识别文件编码：{path}")


def set_run_fonts(run, western: str, east_asia: str, size: float, bold: bool = False) -> None:
    run.font.name = western
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), western)
    rfonts.set(qn("w:hAnsi"), western)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), western)


def add_no_proof(run) -> None:
    rpr = run._element.get_or_add_rPr()
    if rpr.find(qn("w:noProof")) is None:
        rpr.append(OxmlElement("w:noProof"))


def set_cell_text_style(cell, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_fonts(run, "Times New Roman", "SimSun", 8.5, bold=bold)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def ensure_code_style(document: Document):
    style_name = "Appendix Code"
    if style_name in [style.name for style in document.styles]:
        return document.styles[style_name]
    style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(6.5)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = Pt(8.0)
    style.paragraph_format.left_indent = Pt(3)
    style.paragraph_format.right_indent = Pt(0)
    style.paragraph_format.keep_together = False
    style.paragraph_format.keep_with_next = False
    style.paragraph_format.widow_control = False
    return style


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_normal_paragraph(document: Document, text: str, bold_prefix: str | None = None):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        run1 = paragraph.add_run(bold_prefix)
        set_run_fonts(run1, "Times New Roman", "SimSun", 10.5, bold=True)
        run2 = paragraph.add_run(text[len(bold_prefix):])
        set_run_fonts(run2, "Times New Roman", "SimSun", 10.5)
    else:
        run = paragraph.add_run(text)
        set_run_fonts(run, "Times New Roman", "SimSun", 10.5)
    return paragraph


def add_environment_table(document: Document) -> None:
    rows = [
        ("编程语言", "Python 3.12.13（64位）", "三问优化、结果核验和论文绘图"),
        ("数值计算", "NumPy 2.5.2", "数组运算、统计计算与结果校验"),
        ("优化求解", "SciPy 1.18.0（milp/HiGHS）", "混合整数线性规划与稀疏约束矩阵"),
        ("Excel读写", "openpyxl 3.1.5", "读取附件、输出排班工作簿与人员月历"),
        ("数据整理", "pandas 3.0.1", "绘图数据合并与一致性检查"),
        ("图形绘制", "Matplotlib 3.10.5", "生成PNG、SVG、PDF和TIFF论文图"),
        ("运行平台", "Windows 64位", "源码同样可在配置上述依赖的Linux/macOS环境运行"),
    ]
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(3.0), Cm(5.0), Cm(8.0)]
    headers = ["项目", "版本/配置", "用途"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        cell.width = widths[index]
        shade_cell(cell, "D9EAF7")
        set_cell_margins(cell)
        set_cell_text_style(cell, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for item in rows:
        cells = table.add_row().cells
        for index, text in enumerate(item):
            cells[index].text = text
            cells[index].width = widths[index]
            set_cell_margins(cells[index])
            set_cell_text_style(
                cells[index],
                bold=False,
                align=WD_ALIGN_PARAGRAPH.CENTER if index < 2 else WD_ALIGN_PARAGRAPH.LEFT,
            )


def add_source_inventory(document: Document, sources: list[dict]) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(1.6), Cm(5.2), Cm(9.2)]
    headers = ["编号", "程序文件", "主要功能"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        cell.width = widths[index]
        shade_cell(cell, "D9EAF7")
        set_cell_margins(cell)
        set_cell_text_style(cell, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for source in sources:
        cells = table.add_row().cells
        values = (source["number"], source["file"].name, source["purpose"])
        for index, text in enumerate(values):
            cells[index].text = str(text)
            cells[index].width = widths[index]
            set_cell_margins(cells[index])
            set_cell_text_style(
                cells[index],
                bold=False,
                align=WD_ALIGN_PARAGRAPH.CENTER if index < 2 else WD_ALIGN_PARAGRAPH.LEFT,
            )


def add_code_listing(
    document: Document,
    source: dict,
    code_style,
    page_break_before: bool = True,
) -> None:
    heading = document.add_paragraph(style="Heading 3")
    heading.paragraph_format.page_break_before = page_break_before
    heading.paragraph_format.keep_with_next = True
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading.add_run(f"附录{source['number']} {source['title']}（{source['file'].name}）")

    meta = document.add_paragraph(style="Normal")
    meta.paragraph_format.first_line_indent = Cm(0)
    meta.paragraph_format.space_after = Pt(3)
    meta.paragraph_format.keep_with_next = True
    run = meta.add_run("程序语言与版本：Python 3.12.13。功能：")
    set_run_fonts(run, "Times New Roman", "SimSun", 9.0, bold=True)
    run = meta.add_run(source["purpose"])
    set_run_fonts(run, "Times New Roman", "SimSun", 9.0)

    code = read_text_safely(source["file"]).replace("\r\n", "\n").replace("\r", "\n")
    lines = code.split("\n")
    block_size = 60
    for start in range(0, len(lines), block_size):
        block = "\n".join(lines[start : start + block_size])
        paragraph = document.add_paragraph(style=code_style)
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.widow_control = False
        p_pr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F7F7F7")
        p_pr.append(shd)
        run = paragraph.add_run(block)
        set_run_fonts(run, "Consolas", "Microsoft YaHei", 6.5)
        add_no_proof(run)


def validate_sources(sources: list[dict]) -> None:
    missing = [str(item["file"]) for item in sources if not item["file"].is_file()]
    if missing:
        raise FileNotFoundError("缺少附录源码文件：\n" + "\n".join(missing))


def create_decimal_numbering(document: Document) -> int:
    """创建从1开始的独立十进制编号，避免沿用正文列表序号。"""
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract_num.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.append(p_pr)
    abstract_num.append(level)
    numbering.append(abstract_num)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def main() -> None:
    all_sources = MODEL_SOURCES + FIGURE_SOURCES
    validate_sources(all_sources)
    if not INPUT_DOCX.is_file():
        raise FileNotFoundError(f"未找到论文终稿：{INPUT_DOCX}")

    document = Document(INPUT_DOCX)
    code_style = ensure_code_style(document)
    set_update_fields(document)

    attachment_heading = None
    for paragraph in reversed(document.paragraphs):
        if paragraph.text.strip() == "附件":
            attachment_heading = paragraph
            break
    if attachment_heading is None:
        attachment_heading = document.add_paragraph(style="Heading 1")
    attachment_heading.text = "附录"
    attachment_heading.style = document.styles["Heading 1"]
    attachment_heading.paragraph_format.page_break_before = True

    next_paragraph = None
    paragraphs = document.paragraphs
    attachment_index = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph._p is attachment_heading._p
    )
    if attachment_index + 1 < len(paragraphs):
        candidate = paragraphs[attachment_index + 1]
        if not candidate.text.strip():
            next_paragraph = candidate
    if next_paragraph is None:
        next_paragraph = document.add_paragraph(style="Heading 2")
    next_paragraph.text = "附录A 程序运行环境与复现说明"
    next_paragraph.style = document.styles["Heading 2"]

    add_normal_paragraph(
        document,
        "本附录给出论文三个问题的完整求解程序、结果后处理程序、独立检验程序和全部绘图程序。程序中保留了模型变量定义、约束构建、词典序优化、整数流分解、结果核验及图形输出等关键步骤的注释，可用于复现正文中的主要数值结果。",
    )
    add_environment_table(document)

    document.add_heading("附录A.1 建议运行顺序", level=3)
    steps = [
        ("步骤一：", "将原始数据文件“第一轮B题附件.xlsx”与求解程序放在同一项目目录中。"),
        ("步骤二：", "依次运行solve_question1.py和solve_question2.py，得到问题一和问题二的每日班次与人数结果。"),
        ("步骤三：", "运行solve_question3.py；需要融合日内细化结果时，再运行solve_question3_hybrid.py和optimize_question3_results.py。"),
        ("步骤四：", "运行verify_question3.py对具名员工月历和每日覆盖进行独立核验，再运行validate_all_models.py完成三问有效性与鲁棒性检验。"),
        ("步骤五：", "最后运行附录C中的绘图程序。对含有本机绝对输出路径的脚本，应先将out或savefig路径改为当前项目目录。"),
    ]
    for label, text in steps:
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(label)
        set_run_fonts(run, "Times New Roman", "SimSun", 10.0, bold=True)
        run = paragraph.add_run(text)
        set_run_fonts(run, "Times New Roman", "SimSun", 10.0)

    document.add_heading("附录A.2 程序文件清单", level=3)
    add_source_inventory(document, all_sources)

    heading_b = document.add_paragraph(style="Heading 2")
    heading_b.paragraph_format.page_break_before = True
    heading_b.add_run("附录B 三问模型求解、后处理与检验程序")
    add_normal_paragraph(
        document,
        "本部分代码对应正文问题一至问题三的模型建立、混合整数规划求解、具名排班分解及模型有效性和鲁棒性检验。",
    )
    for index, source in enumerate(MODEL_SOURCES):
        add_code_listing(document, source, code_style, page_break_before=index > 0)

    heading_c = document.add_paragraph(style="Heading 2")
    heading_c.paragraph_format.page_break_before = True
    heading_c.add_run("附录C 论文插图绘制程序")
    add_normal_paragraph(
        document,
        "本部分列出论文中人员规模、库存变化、连续工日、到货量热力图、模型结构图和敏感性图的完整绘图源码。",
    )
    for index, source in enumerate(FIGURE_SOURCES):
        add_code_listing(document, source, code_style, page_break_before=index > 0)

    document.save(OUTPUT_DOCX)
    print(f"已生成：{OUTPUT_DOCX}")
    print(f"已嵌入源码文件：{len(all_sources)}个")


if __name__ == "__main__":
    main()
