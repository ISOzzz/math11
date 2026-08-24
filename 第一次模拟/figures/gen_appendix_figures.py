#!/usr/bin/env python3
"""Generate publication-quality supplementary figures from the saved results.

Environment used for reproduction:
Python 3.12, matplotlib 3.10.5, openpyxl 3.1.5.
The script is read-only with respect to the source workbooks.
"""

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.ticker import MaxNLocator
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "运行结果_20260819"
OUTPUT = Path(__file__).resolve().parent

Q1_FILE = RESULTS / "问题一优化结果_人数均衡.xlsx"
Q2_FILE = RESULTS / "问题二优化结果.xlsx"
Q3_FILE = RESULTS / "问题三联合优化结果.xlsx"

# Colorblind-safe Okabe-Ito palette.
BLUE = "#0072B2"
SKY = "#56B4E9"
GREEN = "#009E73"
ORANGE = "#E69F00"
VERMILION = "#D55E00"
GRAY = "#7A7A7A"

plt.rcParams.update(
    {
        "font.family": "sans-serif",
        "font.sans-serif": ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "font.size": 10,
        "axes.titlesize": 11,
        "axes.titleweight": "bold",
        "axes.labelsize": 10,
        "legend.fontsize": 8.5,
        "legend.frameon": False,
        "figure.dpi": 300,
        "savefig.dpi": 300,
        "savefig.bbox": "tight",
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.grid": True,
        "grid.alpha": 0.18,
        "grid.linestyle": "--",
        "lines.linewidth": 1.8,
        "lines.markersize": 4.5,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    }
)


def numeric_rows(path: Path, sheet_index: int, columns: tuple[int, ...]) -> np.ndarray:
    """Read selected one-based columns from a worksheet, excluding its header."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[sheet_index]
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = [row[column - 1] for column in columns]
        if values[0] is None:
            continue
        rows.append([float(value) for value in values])
    workbook.close()
    return np.asarray(rows, dtype=float)


def save_figure(fig: plt.Figure, stem: str) -> None:
    fig.savefig(OUTPUT / f"{stem}.pdf")
    fig.savefig(OUTPUT / f"{stem}.png", dpi=300)
    plt.close(fig)


def plot_daily_staffing() -> None:
    q1 = numeric_rows(Q1_FILE, 0, (1, 4))
    q2 = numeric_rows(Q2_FILE, 0, (1, 5))
    q3 = numeric_rows(Q3_FILE, 1, (1, 3))

    fig, ax = plt.subplots(figsize=(7.0, 3.35))
    ax.plot(q1[:, 0], q1[:, 1], color=BLUE, marker="o", markevery=3, label="问题一：最低日用工")
    ax.plot(q2[:, 0], q2[:, 1], color=ORANGE, marker="s", markevery=3, label="问题二：时限下最低需求")
    ax.plot(q3[:, 0], q3[:, 1], color=GREEN, marker="^", markevery=3, label="问题三：实际出勤")
    ax.set_title("三问每日人员规模对比")
    ax.set_xlabel("日期/天")
    ax.set_ylabel("人数/人")
    ax.set_xlim(1, 30)
    ax.xaxis.set_major_locator(MaxNLocator(integer=True, nbins=10))
    ax.legend(ncol=3, loc="upper center", bbox_to_anchor=(0.5, 1.00))
    save_figure(fig, "fig_a_daily_staffing")


def plot_hourly_inventory() -> None:
    q1 = numeric_rows(Q1_FILE, 2, (1, 2, 5))
    q2 = numeric_rows(Q2_FILE, 3, (1, 2, 9))
    hours = np.arange(24)
    mean_q1 = np.asarray([q1[q1[:, 1] == hour, 2].mean() for hour in hours])
    mean_q2 = np.asarray([q2[q2[:, 1] == hour, 2].mean() for hour in hours])

    fig, ax = plt.subplots(figsize=(6.7, 3.35))
    ax.plot(hours, mean_q1, color=BLUE, marker="o", label="问题一平均库存")
    ax.plot(hours, mean_q2, color=VERMILION, marker="s", label="问题二平均库存")
    ax.fill_between(hours, mean_q1, mean_q2, color=SKY, alpha=0.13, label="库存差异")
    ax.axvline(15, color=GRAY, linestyle=":", linewidth=1.3)
    ax.text(15.2, max(mean_q1.max(), mean_q2.max()) * 0.90, "16:00截止", color=GRAY, fontsize=8)
    ax.set_title("两种日排班模型的逐小时平均库存")
    ax.set_xlabel("小时结束时刻")
    ax.set_ylabel("平均未处理货量/件")
    ax.set_xticks(np.arange(0, 24, 2))
    ax.legend(loc="upper left")
    save_figure(fig, "fig_b_hourly_inventory")


def plot_q3_surplus() -> None:
    data = numeric_rows(Q3_FILE, 1, (1, 2, 3))
    day, minimum, attendance = data.T

    fig, ax = plt.subplots(figsize=(7.0, 3.35))
    ax.plot(day, minimum, color=BLUE, marker="o", markevery=3, label="问题二最低需求")
    ax.plot(day, attendance, color=GREEN, marker="^", markevery=3, label="问题三实际出勤")
    ax.fill_between(day, minimum, attendance, color=ORANGE, alpha=0.28, label="制度性冗余")
    ax.set_title("连续工日约束下每日需求与实际出勤")
    ax.set_xlabel("日期/天")
    ax.set_ylabel("人数/人")
    ax.set_xlim(1, 30)
    ax.xaxis.set_major_locator(MaxNLocator(integer=True, nbins=10))
    ax.legend(ncol=3, loc="upper center", bbox_to_anchor=(0.5, 1.00))
    save_figure(fig, "fig_c_q3_surplus")


def plot_streak_distribution() -> None:
    data = numeric_rows(Q3_FILE, 5, (1, 4))
    streaks, counts = np.unique(data[:, 1].astype(int), return_counts=True)

    fig, ax = plt.subplots(figsize=(4.6, 3.2))
    colors = [SKY if streak < 7 else VERMILION for streak in streaks]
    bars = ax.bar(streaks.astype(str), counts, color=colors, width=0.58, edgecolor="white")
    for bar, value in zip(bars, counts):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + 7,
            f"{value}人\n({value / counts.sum():.2%})",
            ha="center",
            va="bottom",
            fontsize=9,
        )
    ax.set_title("员工最大连续工作天数分布")
    ax.set_xlabel("最大连续工作天数/天")
    ax.set_ylabel("员工人数/人")
    ax.set_ylim(0, max(counts) * 1.20)
    ax.grid(axis="x", visible=False)
    save_figure(fig, "fig_d_streak_distribution")


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    plot_daily_staffing()
    plot_hourly_inventory()
    plot_q3_surplus()
    plot_streak_distribution()
    print("Generated four appendix figures in PNG and PDF formats.")


if __name__ == "__main__":
    main()
