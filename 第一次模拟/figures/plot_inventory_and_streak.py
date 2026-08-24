#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""根据三问求解结果生成两张论文图：

1. 两种日排班模型的逐小时平均库存；
2. 员工最大连续工作天数分布。

依赖：Python 3.10+、matplotlib、numpy、openpyxl。
源工作簿只读，脚本不会修改求解结果。

运行：
    python plot_inventory_and_streak.py
    python plot_inventory_and_streak.py --data-dir "运行结果目录" --output-dir "图片目录"
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
from collections import Counter
from pathlib import Path


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent

# 兼容项目内置依赖目录；在普通 Python 环境中会直接使用已安装的包。
LOCAL_DEPS = PROJECT_ROOT / ".runtime_deps"
if LOCAL_DEPS.is_dir():
    sys.path.insert(0, str(LOCAL_DEPS))

# 在服务器或无图形界面的命令行环境中也能正常出图。
os.environ.setdefault("MPLBACKEND", "Agg")
os.environ.setdefault(
    "MPLCONFIGDIR",
    str(Path(tempfile.gettempdir()) / "inventory_streak_mplconfig"),
)

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from openpyxl import load_workbook


MM_PER_INCH = 25.4
WIDTH_MM = 183.0
SINGLE_COLUMN_WIDTH_MM = 89.0

# 色盲友好配色，并通过线型/标记提供非颜色区分。
BLUE = "#0072B2"
SKY = "#56B4E9"
VERMILION = "#D55E00"
GRAY = "#707780"
GRID = "#D7DDE5"
TEXT = "#20252B"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="绘制平均库存图和连续工作天数分布图")
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=PROJECT_ROOT / "运行结果_20260819",
        help="包含三问 Excel 求解结果的目录",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=SCRIPT_DIR,
        help="图片输出目录",
    )
    return parser.parse_args()


def choose_chinese_font() -> str:
    available = {item.name for item in font_manager.fontManager.ttflist}
    candidates = [
        "Microsoft YaHei",
        "Microsoft YaHei UI",
        "SimHei",
        "Noto Sans CJK SC",
        "Source Han Sans SC",
        "Arial Unicode MS",
    ]
    for family in candidates:
        if family in available:
            return family
    print("警告：未检测到常用中文字体，请安装微软雅黑或思源黑体。")
    return "DejaVu Sans"


def configure_style() -> str:
    chinese_font = choose_chinese_font()
    mpl.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": [chinese_font, "Arial", "DejaVu Sans"],
            "font.size": 8.0,
            "axes.titlesize": 10.5,
            "axes.titleweight": "semibold",
            "axes.labelsize": 8.5,
            "xtick.labelsize": 7.5,
            "ytick.labelsize": 7.5,
            "legend.fontsize": 7.5,
            "axes.unicode_minus": False,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.linewidth": 0.8,
            "axes.edgecolor": GRAY,
            "axes.labelcolor": TEXT,
            "xtick.color": TEXT,
            "ytick.color": TEXT,
            "text.color": TEXT,
            "legend.frameon": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
            "savefig.dpi": 600,
        }
    )
    return chinese_font


def read_columns_by_header(
    path: Path,
    sheet_name: str,
    required_headers: tuple[str, ...],
) -> dict[str, np.ndarray]:
    """按表头名称读取数据，避免 Excel 列顺序调整导致读错列。"""
    if not path.is_file():
        raise FileNotFoundError(f"未找到数据文件：{path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise KeyError(f"{path.name}中不存在工作表‘{sheet_name}’。")
        sheet = workbook[sheet_name]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        header_to_index = {str(value): index for index, value in enumerate(header_row) if value is not None}

        missing = [header for header in required_headers if header not in header_to_index]
        if missing:
            raise KeyError(f"{path.name}/{sheet_name}缺少列：{missing}")

        values: dict[str, list[float]] = {header: [] for header in required_headers}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            first_value = row[header_to_index[required_headers[0]]]
            if first_value is None:
                continue
            for header in required_headers:
                cell_value = row[header_to_index[header]]
                if cell_value is None:
                    raise ValueError(f"{path.name}/{sheet_name}的‘{header}’列存在缺失值。")
                values[header].append(float(cell_value))
    finally:
        workbook.close()

    arrays = {header: np.asarray(column, dtype=float) for header, column in values.items()}
    if not arrays[required_headers[0]].size:
        raise ValueError(f"{path.name}/{sheet_name}没有可用数据。")
    if not all(np.isfinite(column).all() for column in arrays.values()):
        raise ValueError(f"{path.name}/{sheet_name}存在非有限数值。")
    return arrays


def load_hourly_inventory(data_dir: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    q1 = read_columns_by_header(
        data_dir / "问题一优化结果_人数均衡.xlsx",
        "逐小时详情",
        ("天", "小时", "剩余货量"),
    )
    q2 = read_columns_by_header(
        data_dir / "问题二优化结果.xlsx",
        "逐小时详情",
        ("天", "小时", "总剩余"),
    )

    hours = np.arange(24, dtype=int)
    expected_days = set(range(1, 31))
    mean_q1 = np.empty(24, dtype=float)
    mean_q2 = np.empty(24, dtype=float)

    for label, data in (("问题一", q1), ("问题二", q2)):
        days = data["天"].astype(int)
        hour_values = data["小时"].astype(int)
        if set(days) != expected_days:
            raise ValueError(f"{label}逐小时数据必须完整覆盖第1—30天。")
        if set(hour_values) != set(hours):
            raise ValueError(f"{label}逐小时数据必须完整覆盖0—23时。")
        for hour in hours:
            if np.count_nonzero(hour_values == hour) != 30:
                raise ValueError(f"{label}的{hour}时应有30个每日观测值。")

    for index, hour in enumerate(hours):
        mean_q1[index] = q1["剩余货量"][q1["小时"].astype(int) == hour].mean()
        mean_q2[index] = q2["总剩余"][q2["小时"].astype(int) == hour].mean()

    return hours, mean_q1, mean_q2


def load_streak_distribution(data_dir: Path) -> tuple[np.ndarray, np.ndarray]:
    employee = read_columns_by_header(
        data_dir / "问题三联合优化结果.xlsx",
        "人员检验",
        ("工人编号", "最大连续工作天数"),
    )
    worker_ids = employee["工人编号"].astype(int)
    streak_values_float = employee["最大连续工作天数"]
    if len(worker_ids) != len(set(worker_ids)):
        raise ValueError("人员检验表存在重复工人编号。")
    if not np.allclose(streak_values_float, np.rint(streak_values_float)):
        raise ValueError("最大连续工作天数必须为整数。")

    streak_values = streak_values_float.astype(int)
    if (streak_values > 7).any():
        raise ValueError("检测到员工最大连续工作天数超过7天。")
    count_map = Counter(streak_values.tolist())
    streaks = np.asarray(sorted(count_map), dtype=int)
    counts = np.asarray([count_map[value] for value in streaks], dtype=int)
    return streaks, counts


def style_axes(ax: plt.Axes, grid_axis: str = "y") -> None:
    ax.grid(axis=grid_axis, color=GRID, linewidth=0.55, linestyle="--", alpha=0.65)
    ax.set_axisbelow(True)
    ax.tick_params(axis="both", direction="out", length=3.0, width=0.7)


def plot_hourly_inventory(
    hours: np.ndarray,
    mean_q1: np.ndarray,
    mean_q2: np.ndarray,
) -> plt.Figure:
    width_mm, height_mm = WIDTH_MM, 92.0
    fig, ax = plt.subplots(
        figsize=(width_mm / MM_PER_INCH, height_mm / MM_PER_INCH),
        layout="constrained",
    )

    line_q1, = ax.plot(
        hours,
        mean_q1,
        color=BLUE,
        linewidth=1.7,
        marker="o",
        markersize=3.2,
        markeredgewidth=0,
        label="问题一平均库存",
        zorder=3,
    )
    line_q2, = ax.plot(
        hours,
        mean_q2,
        color=VERMILION,
        linewidth=1.7,
        marker="s",
        markersize=3.1,
        markeredgewidth=0,
        label="问题二平均库存",
        zorder=4,
    )
    difference = ax.fill_between(
        hours,
        mean_q1,
        mean_q2,
        color=SKY,
        alpha=0.15,
        linewidth=0,
        label="库存差异",
        zorder=1,
    )

    cutoff_hour = 15
    ax.axvline(cutoff_hour, color=GRAY, linestyle=":", linewidth=1.15, zorder=2)
    annotation_y = max(mean_q1.max(), mean_q2.max()) * 0.91
    ax.annotate(
        "16:00截止",
        xy=(cutoff_hour, annotation_y),
        xytext=(5, 0),
        textcoords="offset points",
        ha="left",
        va="center",
        fontsize=7.5,
        color=GRAY,
    )

    ax.set_title("两种日排班模型的逐小时平均库存", pad=7)
    ax.set_xlabel("小时结束时刻")
    ax.set_ylabel("平均未处理货量/件")
    ax.set_xlim(-0.8, 23.2)
    ax.set_xticks(np.arange(0, 24, 2))
    style_axes(ax, grid_axis="both")
    ax.legend(
        handles=[line_q1, line_q2, difference],
        loc="upper left",
        handlelength=2.2,
        handletextpad=0.6,
        borderaxespad=0.5,
    )
    return fig


def plot_streak_distribution(streaks: np.ndarray, counts: np.ndarray) -> plt.Figure:
    width_mm, height_mm = SINGLE_COLUMN_WIDTH_MM, 78.0
    fig, ax = plt.subplots(
        figsize=(width_mm / MM_PER_INCH, height_mm / MM_PER_INCH),
        layout="constrained",
    )

    colors = [SKY if streak < 7 else VERMILION for streak in streaks]
    x_positions = np.arange(len(streaks))
    bars = ax.bar(
        x_positions,
        counts,
        color=colors,
        width=0.58,
        edgecolor="white",
        linewidth=0.7,
    )

    total = int(counts.sum())
    label_offset = max(counts) * 0.025
    for bar, count in zip(bars, counts):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + label_offset,
            f"{count}人\n({count / total:.2%})",
            ha="center",
            va="bottom",
            fontsize=7.5,
        )

    ax.set_title("员工最大连续工作天数分布", pad=7)
    ax.set_xlabel("最大连续工作天数/天")
    ax.set_ylabel("员工人数/人")
    ax.set_xticks(x_positions, [str(value) for value in streaks])
    ax.set_ylim(0, max(counts) * 1.19)
    style_axes(ax, grid_axis="y")
    return fig


def save_figure(fig: plt.Figure, output_stem: Path) -> None:
    """输出600 dpi位图和文字可编辑的矢量图。"""
    fig.savefig(output_stem.with_suffix(".png"), dpi=600, facecolor="white")
    fig.savefig(output_stem.with_suffix(".tiff"), dpi=600, facecolor="white")
    fig.savefig(output_stem.with_suffix(".svg"), facecolor="white")
    fig.savefig(output_stem.with_suffix(".pdf"), facecolor="white")
    plt.close(fig)


def main() -> None:
    args = parse_args()
    data_dir = args.data_dir.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    font_name = configure_style()

    hours, mean_q1, mean_q2 = load_hourly_inventory(data_dir)
    streaks, counts = load_streak_distribution(data_dir)

    inventory_figure = plot_hourly_inventory(hours, mean_q1, mean_q2)
    save_figure(inventory_figure, output_dir / "图_两种日排班模型逐小时平均库存")

    streak_figure = plot_streak_distribution(streaks, counts)
    save_figure(streak_figure, output_dir / "图_员工最大连续工作天数分布")

    print(f"中文字体：{font_name}")
    print(f"平均库存数据：{len(hours)}个小时，每小时30个每日观测值。")
    print(f"连续工日分布：{dict(zip(streaks.tolist(), counts.tolist()))}，共{counts.sum()}人。")
    print(f"图片已输出至：{output_dir}")


if __name__ == "__main__":
    main()
