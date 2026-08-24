from pathlib import Path
import re
import zipfile

from docx import Document
from lxml import etree


path = Path(r"D:\第一次模拟\运行结果_20260819\数学建模论文终稿_模块化摘要.docx")
doc = Document(path)

with zipfile.ZipFile(path) as package:
    zip_error = package.testzip()
    root = etree.fromstring(package.read("word/document.xml"))

ns = {
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}

paragraph_text = "\n".join(p.text for p in doc.paragraphs)
table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
all_text = paragraph_text + "\n" + table_text

expected_equations = (
    [f"（5.{i}）" for i in range(1, 16)]
    + [f"（6.{i}）" for i in range(1, 13)]
    + [f"（7.{i}）" for i in range(1, 21)]
    + ["（8.1）"]
)
missing_equations = [label for label in expected_equations if label not in paragraph_text]
found_equations = re.findall(r"（[5-8]\.\d+）", paragraph_text)

expected_captions = [
    "表4-1 主要符号说明",
    "表5-1 问题一核心求解结果",
    "表6-1 问题二核心求解结果",
    "表7-1 问题三核心求解结果",
    "图6-1 问题一与问题二逐小时平均库存对比",
    "图7-1 第三问每日最低需求与实际出勤",
    "图7-2 员工最大连续工作天数分布",
    "表9-1 三问核心指标综合比较",
    "图9-1 三问每日人员规模对比",
    "表9-2 处理效率敏感性分析",
    "表9-3 模型有效性检验汇总",
    "表9-4 随机扰动鲁棒性检验汇总",
]
missing_captions = [caption for caption in expected_captions if caption not in paragraph_text]

required_values = [
    "10791",
    "359.70",
    "99.88%",
    "11858",
    "395.27",
    "98.26%",
    "581",
    "13363",
    "53",
    "0.473%",
    "1.240%",
    "1.721%",
]
missing_values = [value for value in required_values if value not in all_text]

forbidden_fragments = [
    "建议在论文中",
    "建议绘制",
    "目录将在打开文档时自动更新",
    "图1 ",
    "图2 ",
    "图3 ",
    "图4 ",
    "表1 ",
    "表2 ",
    "表3 ",
    "表4 ",
    "\x0b",
]
found_forbidden = [fragment for fragment in forbidden_fragments if fragment in all_text]

print(f"ZIP_TEST={zip_error}")
print(f"PARAGRAPHS={len(doc.paragraphs)}")
print(f"HEADINGS={sum(p.style.name.startswith('Heading ') for p in doc.paragraphs)}")
print(f"TABLES={len(doc.tables)}")
print(f"DRAWINGS={len(root.xpath('.//w:drawing', namespaces=ns))}")
print(f"OMATH_NODES={len(root.xpath('.//m:oMath', namespaces=ns))}")
print(f"SECTIONS={len(doc.sections)}")
print(f"EQUATION_LABELS={len(found_equations)}")
print(f"MISSING_EQUATIONS={missing_equations}")
print(f"MISSING_CAPTIONS={missing_captions}")
print(f"MISSING_VALUES={missing_values}")
print(f"FORBIDDEN_FRAGMENTS={found_forbidden}")

assert zip_error is None
assert len(doc.tables) == 8
assert len(root.xpath('.//w:drawing', namespaces=ns)) == 4
assert len(doc.sections) == 1
assert len(set(found_equations)) == 48 and not missing_equations
assert not missing_captions
assert not missing_values
assert not found_forbidden
