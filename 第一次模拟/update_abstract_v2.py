from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"D:\第一次模拟\运行结果_20260819\数学建模论文终稿.docx")
OUTPUT = Path(r"D:\第一次模拟\运行结果_20260819\数学建模论文终稿_摘要二次优化.docx")

ABSTRACT = (
    "随着电子商务与即时配送业务持续增长，物流分拣中心面临到货量波动、处理时限收紧与用工成本上升的多重压力。"
    "针对已知未来30天逐小时到货量、每日可设置5个8小时班次的场景，本文研究货物不得提前处理、早期货物须在16:00前清零、"
    "全部货物须当日完成、工人存在低效小时，以及每名员工月工作23天且连续工作不超过7天等约束下的班次安排与人员配置问题。"
    "按照“日内处理—低效调度—月度配置”的递进逻辑，分别建立库存平衡班次优化模型、时限与低效约束排班模型和连续工日约束人员配置优化模型。"
    "问题一以班次启用、班次人数、小时处理量和库存量为决策变量，通过库存递推禁止提前处理，并采用词典序目标依次最小化总人日、在岗峰值、"
    "小时人员偏差和累计积压；问题二将货物划分为早期货物与普通货物，引入低效小时人数变量，显式刻画效率降至10件/小时造成的能力损失，"
    "同时约束两类库存及截止时间；问题三利用“累计工日—连续工日”状态流网络求得最少招聘人数，并将聚合流分解为具名员工月历。"
    "三问均转化为混合整数线性规划并求解。结果显示，问题一全月需10791人日，日均359.70人，处理能力利用率达99.88%，"
    "30天均达到每人每日处理200件对应的理论下界；问题二需11858人日，较问题一增加9.89%，利用率为98.26%，早期货物均在16:00前清零且全部货物均在24:00前完成；"
    "问题三最少招聘581人，共安排13363人日，每名员工恰工作23天、连续工作不超过7天，全月小时在岗峰值为295人。"
    "模型检验表明，三问最大约束残差与整数性误差均为0，问题一、二逐日独立重求解一致率均为100%，问题三理论下界与可行解同为581人，最优性间隙为0。"
    "对720小时到货量进行20次独立±10%随机扰动后，三问目标值最大绝对变化率分别为0.473%、1.240%和1.721%；"
    "全量同比变化±10%时，人员目标约变化±10%，说明模型能抵抗零均值小时噪声，但对系统性业务增长保持合理敏感。"
    "本文创新在于构建覆盖“班次—低效小时—员工月历”的多层递进优化框架，显式量化低效能力损失，并以状态流分解实现从聚合配置到个体排班的可执行转化。"
    "研究结果可为物流分拣中心的班次设计、人员招聘与轮休管理提供可靠的定量决策依据。"
)


def main() -> None:
    if not 800 <= len(ABSTRACT) <= 1000:
        raise ValueError(f"摘要字符数不符合要求：{len(ABSTRACT)}")

    document = Document(SOURCE)
    if len(document.paragraphs) < 4 or document.paragraphs[1].text.strip() != "摘要":
        raise RuntimeError("未能定位终稿摘要段落")

    paragraph = document.paragraphs[2]
    paragraph.clear()
    run = paragraph.add_run(ABSTRACT)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")

    document.save(OUTPUT)

    check = Document(OUTPUT)
    saved = check.paragraphs[2].text
    if saved != ABSTRACT:
        raise RuntimeError("摘要写入后校验失败")
    print(f"OUTPUT={OUTPUT}")
    print(f"ABSTRACT_CHARS={len(saved)}")
    print(f"ABSTRACT_RUNS={len(check.paragraphs[2].runs)}")


if __name__ == "__main__":
    main()
