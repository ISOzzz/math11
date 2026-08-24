#!/usr/bin/env python3
"""Build the integrated Chinese mathematical-modeling paper draft."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import etree
from latex2mathml.converter import convert as latex_to_mathml


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "运行结果_20260819"
OUTPUT = RESULTS / "数学建模论文初稿_约25页.docx"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

TITLE_SOURCE = RESULTS / "论文题目摘要关键词.md"
PRELIM_SOURCE = RESULTS / "论文问题重述分析假设符号.md"
MODEL_SOURCE = RESULTS / "论文模型建立求解结果分析检验.md"

FIGURES = {
    "daily": ROOT / "figures" / "fig_a_daily_staffing.png",
    "inventory": ROOT / "figures" / "fig_b_hourly_inventory.png",
    "surplus": ROOT / "figures" / "fig_c_q3_surplus.png",
    "streak": ROOT / "figures" / "fig_d_streak_distribution.png",
}

TABLE_TITLES = [
    "表1 主要符号说明",
    "表2 问题一核心求解结果",
    "表3 问题二核心求解结果",
    "表4 问题三核心求解结果",
    "表5 三问核心指标综合比较",
    "表6 处理效率敏感性分析",
    "表7 模型有效性检验汇总",
    "表8 随机扰动鲁棒性检验汇总",
]


CONCLUSION = r"""
# 9 结论与建议

## 9.1 主要结论

本文围绕物流分拣中心的日内班次、时限服务和月度人员配置建立了三个递进的混合整数规划模型。三个模型分别回答“当天最少需要多少人”“考虑16:00截止与低效小时后如何排班”以及“满足固定工日和连续工作限制时最少招聘多少人”。综合结果如下表所示。

| 比较维度 | 问题一 | 问题二 | 问题三 |
|---|---:|---:|---:|
| 核心决策 | 日班次与人数 | 时限、低效小时与人数 | 招聘、出勤与具名月历 |
| 月累计人日 | 10791 | 11858 | 13363 |
| 日均出勤/人 | 359.70 | 395.27 | 445.43 |
| 最大日人数/人 | 537 | 581 | 581 |
| 总体能力利用率 | 99.88% | 98.26% | 87.20% |
| 关键服务约束 | 24:00清空 | 16:00与24:00清空 | 每人23天且连续工日不超过7天 |

问题一的30天人员数全部达到每人每天200件对应的理论下界，说明在允许日内暂存的条件下，5班次结构能够充分利用人力。问题二以增加1067人日为代价，保证早期货物全部在16:00前处理完毕，并使平均小时库存下降48.35%。问题三进一步证明最少招聘人数为581人：第12天的单日需求给出581人的必要下界，状态流模型又构造出581人的完整可行方案，故最优性间隙为0。

## 9.2 敏感性与稳定性结论

正常处理效率是人员规模的关键驱动因素。效率下降时，高负荷日能力下界会直接上升；效率提高则能降低长期招聘规模。重新求解结果如下表所示。

| 场景 | 问题一日均人数 | 问题二日均人数 | 问题三招聘人数 |
|---|---:|---:|---:|
| 正常效率下降10% | 399.67（+11.11%） | 437.13（+10.59%） | 642（+10.50%） |
| 基准效率 | 359.70 | 395.27 | 581 |
| 正常效率提高10% | 327.10（-9.06%） | 360.67（-8.75%） | 531（-8.61%） |

随机小时扰动和系统性业务增长呈现不同影响机制。零均值小时扰动可在日内和月内部分抵消，三个目标的最大绝对变化率均低于2%；而全部小时同向变化10%时，人员目标近似同比变化。因此，模型适合作为确定性基准排班，但企业仍需为持续性业务增长预留弹性。

## 9.3 有效性与鲁棒性汇总

| 检验项目 | 问题一 | 问题二 | 问题三 |
|---|---:|---:|---:|
| 约束可行率 | 100% | 100% | 100% |
| 最大约束残差 | 0 | 0 | 0 |
| 最大整数性误差 | 0 | 0 | 0 |
| 独立重求解/深度审计 | 100%一致 | 100%一致 | 全部通过 |
| 最优性证书 | 30天命中总量下界 | MILP零间隙并复算约束 | 581人下界与可行解闭合 |

| 指标 | 随机扰动平均变化率 | 标准差 | 最大绝对变化率 |
|---|---:|---:|---:|
| 总到货量 | -0.093% | 0.216% | 0.492% |
| 问题一总人日 | -0.077% | 0.215% | 0.473% |
| 问题二总人日 | 0.297% | 0.459% | 1.240% |
| 问题三招聘人数 | -0.052% | 0.924% | 1.721% |

上述结果表明，模型在给定数据和业务规则内同时具备约束正确性、整数可执行性、独立复现性与扰动稳定性。尤其是问题三通过最大单日下界、总人日下界和连续8日窗口下界共同审计，使招聘结论不仅是软件输出，而且具有可解释的数学证书。

## 9.4 管理建议

第一，企业可将问题二的日最低需求作为短期排班基线，将问题三的581人作为固定员工队伍规模，并通过临时工或跨岗人员吸收超出预测区间的业务增长。第二，问题一利用率接近100%，只适用于到货数据较稳定且缓存空间充足的场景；常态运营更适合保留5%—10%的安全能力。第三，应重点改善正常处理效率和高负荷日期的设备可靠性，因为效率变化会近似同比传导到招聘规模。第四，可利用问题三低需求日形成的316551件未用等价能力安排培训、设备维护或跨区支援，从而把制度性冗余转化为组织韧性。第五，实际部署时应采用滚动优化：每日更新未来数日到货量，并限制对已发布班表的改动幅度，使数学最优方案与现场稳定执行相协调。
"""


class PaperBuilder:
    def __init__(self, font_size: float = 12.0, line_spacing: float = 1.25):
        self.doc = Document()
        self.font_size = font_size
        self.line_spacing = line_spacing
        self.table_count = 0
        self.figure_count = 0
        self.equation_count = 0
        self.math_transform = None
        if MML2OMML.exists():
            self.math_transform = etree.XSLT(etree.parse(str(MML2OMML)))
        self._configure_document()

    @staticmethod
    def _set_run_font(run, east_asia="宋体", western="Times New Roman", size=12, bold=None, italic=None):
        run.font.name = western
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), western)
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), western)
        run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic

    def _configure_document(self):
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.35)
        section.footer_distance = Cm(1.35)

        normal = self.doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        normal.font.size = Pt(self.font_size)
        pf = normal.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0.74)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = self.line_spacing
        pf.widow_control = True

        heading_specs = {
            "Heading 1": (16, 12, 6),
            "Heading 2": (14, 9, 4),
            "Heading 3": (12, 6, 2),
        }
        for style_name, (size, before, after) in heading_specs.items():
            style = self.doc.styles[style_name]
            style.font.name = "Arial"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.first_line_indent = Cm(0)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.line_spacing = 1.0
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True

        # Competition papers are space-constrained; allow major headings to
        # follow the preceding section when enough room remains. The heading's
        # keep-with-next setting still prevents an orphaned chapter title.
        self.doc.styles["Heading 1"].paragraph_format.page_break_before = False

        caption = self.doc.styles["Caption"]
        caption.font.name = "Times New Roman"
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        caption.font.size = Pt(10.5)
        caption.font.color.rgb = RGBColor(0, 0, 0)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.first_line_indent = Cm(0)
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(5)
        caption.paragraph_format.keep_with_next = True

        if "Reference" not in [s.name for s in self.doc.styles]:
            ref_style = self.doc.styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
        else:
            ref_style = self.doc.styles["Reference"]
        ref_style.font.name = "Times New Roman"
        ref_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        ref_style.font.size = Pt(10.5)
        ref_style.paragraph_format.left_indent = Cm(0.74)
        ref_style.paragraph_format.first_line_indent = Cm(-0.74)
        ref_style.paragraph_format.space_after = Pt(3)
        ref_style.paragraph_format.line_spacing = 1.15

        self._configure_lists()
        self._set_header_footer(section)
        settings = self.doc.settings._element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)

    def _configure_lists(self):
        for name in ("List Bullet", "List Number"):
            style = self.doc.styles[name]
            style.font.name = "Times New Roman"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            style.font.size = Pt(self.font_size)
            style.paragraph_format.left_indent = Cm(0.74)
            style.paragraph_format.first_line_indent = Cm(-0.37)
            style.paragraph_format.space_after = Pt(1)
            style.paragraph_format.line_spacing = self.line_spacing

    def _set_header_footer(self, section):
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.space_after = Pt(0)
        run = hp.add_run("物流分拣中心多层排班与人员配置优化")
        self._set_run_font(run, east_asia="宋体", size=9)
        run.font.color.rgb = RGBColor(110, 110, 110)

        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(0)
        run = fp.add_run("— ")
        self._set_run_font(run, east_asia="宋体", size=9)
        self._append_field(fp, "PAGE")
        run = fp.add_run(" —")
        self._set_run_font(run, east_asia="宋体", size=9)

    @staticmethod
    def _append_field(paragraph, instruction: str):
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        begin_run = OxmlElement("w:r")
        begin_run.append(begin)

        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = f" {instruction} "
        instr_run = OxmlElement("w:r")
        instr_run.append(instr)

        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        separate_run = OxmlElement("w:r")
        separate_run.append(separate)

        text_run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "1"
        text_run.append(text)

        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        end_run = OxmlElement("w:r")
        end_run.append(end)

        paragraph._p.extend([begin_run, instr_run, separate_run, text_run, end_run])

    def add_front_matter(self, title: str, abstract: str, keywords: str):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(16)
        r = p.add_run(title)
        self._set_run_font(r, east_asia="黑体", size=22, bold=True)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run("摘要")
        self._set_run_font(r, east_asia="黑体", size=14, bold=True)

        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(5)
        self.add_inline_markup(p, abstract)

        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("关键词：")
        self._set_run_font(r, east_asia="黑体", size=12, bold=True)
        r = p.add_run(keywords)
        self._set_run_font(r, east_asia="宋体", size=12)

        self.doc.add_page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run("目录")
        self._set_run_font(r, east_asia="黑体", size=16, bold=True)
        self.add_toc()
        self.doc.add_page_break()

        # The first body heading follows a manual page break; suppress its own page break.
        self.first_heading = True

    def add_toc(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.2
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "目录将在打开文档时自动更新"
        r.append(t)
        fld.append(r)
        p._p.append(fld)

    def add_heading(self, text: str, level: int):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        if getattr(self, "first_heading", False) and level == 1:
            p.paragraph_format.page_break_before = False
            self.first_heading = False
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(text.strip())
        self._set_run_font(r, east_asia="黑体", western="Arial", size={1: 16, 2: 14, 3: 12}[level], bold=True)
        return p

    def add_paragraph(self, text: str, style=None):
        p = self.doc.add_paragraph(style=style)
        if style in ("List Bullet", "List Number"):
            p.paragraph_format.first_line_indent = Cm(-0.37)
        elif style == "Reference":
            pass
        else:
            p.paragraph_format.first_line_indent = Cm(0.74)
        self.add_inline_markup(p, text.strip())
        return p

    def add_inline_markup(self, paragraph, text: str):
        text = text.replace("`", "")
        token_re = re.compile(r"(\*\*.*?\*\*|\$.*?\$)")
        position = 0
        for match in token_re.finditer(text):
            if match.start() > position:
                run = paragraph.add_run(text[position:match.start()])
                self._set_run_font(run, size=self.font_size)
            token = match.group(0)
            if token.startswith("**"):
                run = paragraph.add_run(token[2:-2])
                self._set_run_font(run, size=self.font_size, bold=True)
            else:
                self.append_math(paragraph, token[1:-1], display=False)
            position = match.end()
        if position < len(text):
            run = paragraph.add_run(text[position:])
            self._set_run_font(run, size=self.font_size)

    def append_math(self, paragraph, latex: str, display: bool):
        latex = latex.strip().rstrip("。；，,")
        if self.math_transform is None:
            run = paragraph.add_run(latex)
            self._set_run_font(run, east_asia="Cambria Math", western="Cambria Math", size=self.font_size)
            return False
        try:
            mathml = latex_to_mathml(latex)
            mathml_tree = etree.fromstring(mathml.encode("utf-8"))
            omml = self.math_transform(mathml_tree).getroot()
            paragraph._p.append(omml)
            return True
        except Exception:
            run = paragraph.add_run(latex)
            self._set_run_font(run, east_asia="Cambria Math", western="Cambria Math", size=self.font_size)
            return False

    def add_display_math(self, latex: str):
        self.equation_count += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_together = True
        self.append_math(p, latex.replace("\n", " "), display=True)
        run = p.add_run(f"    （{self.equation_count}）")
        self._set_run_font(run, east_asia="宋体", size=10.5)

    @staticmethod
    def _set_cell_margins(cell, top=70, start=100, bottom=70, end=100):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcMar = tcPr.first_child_found_in("w:tcMar")
        if tcMar is None:
            tcMar = OxmlElement("w:tcMar")
            tcPr.append(tcMar)
        for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
            node = tcMar.find(qn(f"w:{margin}"))
            if node is None:
                node = OxmlElement(f"w:{margin}")
                tcMar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    @staticmethod
    def _set_table_borders(table):
        tblPr = table._tbl.tblPr
        borders = tblPr.first_child_found_in("w:tblBorders")
        if borders is not None:
            tblPr.remove(borders)
        borders = OxmlElement("w:tblBorders")
        for edge, val, size in (
            ("top", "single", "12"),
            ("left", "nil", "0"),
            ("bottom", "single", "12"),
            ("right", "nil", "0"),
            ("insideH", "nil", "0"),
            ("insideV", "nil", "0"),
        ):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:val"), val)
            node.set(qn("w:sz"), size)
            node.set(qn("w:color"), "000000")
            borders.append(node)
        tblPr.append(borders)

        header_tc_prs = [cell._tc.get_or_add_tcPr() for cell in table.rows[0].cells]
        for tcPr in header_tc_prs:
            tcBorders = tcPr.first_child_found_in("w:tcBorders")
            if tcBorders is None:
                tcBorders = OxmlElement("w:tcBorders")
                tcPr.append(tcBorders)
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "8")
            bottom.set(qn("w:color"), "000000")
            tcBorders.append(bottom)

    @staticmethod
    def _set_repeat_table_header(row):
        trPr = row._tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader")
        tblHeader.set(qn("w:val"), "true")
        trPr.append(tblHeader)

    def add_table(self, rows: list[list[str]]):
        self.table_count += 1
        if self.table_count <= len(TABLE_TITLES):
            cap = self.doc.add_paragraph(TABLE_TITLES[self.table_count - 1], style="Caption")
            cap.paragraph_format.keep_with_next = True

        cols = max(len(row) for row in rows)
        table = self.doc.add_table(rows=len(rows), cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        self._set_table_borders(table)
        self._set_repeat_table_header(table.rows[0])

        total_width = 9070
        if cols == 2:
            widths = [5700, 3370]
        elif cols == 3 and self.table_count == 1:
            widths = [1600, 5700, 1770]
        elif cols == 3:
            widths = [3200, 2600, 3270]
        elif cols == 4:
            widths = [2800, 2090, 2090, 2090]
        elif cols == 6:
            widths = [1750, 1400, 1400, 1400, 1400, 1720]
        else:
            widths = [total_width // cols] * cols
            widths[-1] += total_width - sum(widths)

        tblPr = table._tbl.tblPr
        tblW = tblPr.first_child_found_in("w:tblW")
        tblW.set(qn("w:w"), str(total_width))
        tblW.set(qn("w:type"), "dxa")
        tblLayout = OxmlElement("w:tblLayout")
        tblLayout.set(qn("w:type"), "fixed")
        tblPr.append(tblLayout)

        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for width in widths:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(width))
            grid.append(col)

        for row_index, source_row in enumerate(rows):
            for col_index in range(cols):
                cell = table.rows[row_index].cells[col_index]
                cell.width = Cm(widths[col_index] / 567)
                tcW = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
                tcW.set(qn("w:w"), str(widths[col_index]))
                tcW.set(qn("w:type"), "dxa")
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                self._set_cell_margins(cell)
                text = source_row[col_index].strip() if col_index < len(source_row) else ""
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (cols == 3 and col_index == 1) else WD_ALIGN_PARAGRAPH.CENTER
                self.add_inline_markup(p, text)
                for run in p.runs:
                    self._set_run_font(run, east_asia="宋体", size=9.5, bold=(row_index == 0))
            if row_index == 0:
                for cell in table.rows[row_index].cells:
                    shading = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls("w")))
                    cell._tc.get_or_add_tcPr().append(shading)

        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        self._insert_figure_for_table(self.table_count)

    def _insert_figure_for_table(self, table_number: int):
        if table_number == 3:
            self.add_figure(FIGURES["inventory"], "图1 问题一与问题二逐小时平均库存对比", width_cm=14.2)
        elif table_number == 4:
            self.add_figure(FIGURES["surplus"], "图2 第三问每日最低需求与实际出勤", width_cm=14.2)
            self.add_figure(FIGURES["streak"], "图3 员工最大连续工作天数分布", width_cm=10.0)
        elif table_number == 5:
            self.add_figure(FIGURES["daily"], "图4 三问每日人员规模对比", width_cm=14.2)

    def add_figure(self, path: Path, caption: str, width_cm: float):
        if not path.exists():
            return
        self.figure_count += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        cp = self.doc.add_paragraph(caption, style="Caption")
        cp.paragraph_format.keep_with_next = False

    def add_markdown(self, markdown: str):
        lines = markdown.replace("\r\n", "\n").split("\n")
        i = 0
        paragraph_buffer: list[str] = []

        def flush_paragraph():
            nonlocal paragraph_buffer
            if paragraph_buffer:
                text = "".join(s.strip() for s in paragraph_buffer)
                if text and "建议在论文中" not in text and "建议绘制" not in text:
                    style = "Reference" if re.match(r"^\[\d+\]", text) else None
                    self.add_paragraph(text, style=style)
                paragraph_buffer = []

        while i < len(lines):
            raw = lines[i]
            stripped = raw.strip()
            if not stripped:
                flush_paragraph()
                i += 1
                continue
            if stripped.startswith(">"):
                flush_paragraph()
                i += 1
                continue
            if stripped.startswith("$$"):
                flush_paragraph()
                equation_lines = []
                if stripped != "$$":
                    equation_lines.append(stripped[2:])
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    equation_lines.append(lines[i].strip())
                    i += 1
                if i < len(lines):
                    last = lines[i].strip()
                    if last != "$$":
                        equation_lines.append(last[:-2])
                self.add_display_math(" ".join(equation_lines))
                i += 1
                continue
            if stripped.startswith("#"):
                flush_paragraph()
                match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
                if match:
                    self.add_heading(match.group(2), len(match.group(1)))
                i += 1
                continue
            if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1].strip()):
                flush_paragraph()
                table_lines = [stripped]
                i += 2  # Skip delimiter row.
                while i < len(lines) and lines[i].strip().startswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                parsed_rows = []
                for table_line in table_lines:
                    parsed_rows.append([cell.strip() for cell in table_line.strip("|").split("|")])
                self.add_table(parsed_rows)
                continue
            bullet = re.match(r"^[-*]\s+(.*)$", stripped)
            numbered = re.match(r"^\d+[.)]\s+(.*)$", stripped)
            if bullet:
                flush_paragraph()
                self.add_paragraph(bullet.group(1), style="List Bullet")
                i += 1
                continue
            if numbered:
                flush_paragraph()
                self.add_paragraph(numbered.group(1), style="List Number")
                i += 1
                continue
            paragraph_buffer.append(stripped)
            i += 1
        flush_paragraph()

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.core_properties.title = "基于多层混合整数规划的物流分拣排班问题研究"
        self.doc.core_properties.subject = "数学建模竞赛论文初稿"
        self.doc.core_properties.author = "数学建模参赛队"
        self.doc.core_properties.keywords = "物流分拣；混合整数规划；库存平衡；时间窗；人员配置"
        self.doc.save(path)


def parse_front_matter(text: str):
    title_match = re.search(r"^#\s+(.+)$", text, re.M)
    abstract_match = re.search(r"## 摘要\s*(.*?)\s*## 关键词", text, re.S)
    keywords_match = re.search(r"## 关键词\s*(.*)$", text, re.S)
    if not (title_match and abstract_match and keywords_match):
        raise ValueError("Front matter source is incomplete.")
    abstract = "".join(line.strip() for line in abstract_match.group(1).splitlines() if line.strip())
    keywords = "".join(line.strip() for line in keywords_match.group(1).splitlines() if line.strip())
    return title_match.group(1).strip(), abstract, keywords


def prepare_body() -> str:
    prelim = PRELIM_SOURCE.read_text(encoding="utf-8").strip()
    model = MODEL_SOURCE.read_text(encoding="utf-8")
    start = model.index("# 5 问题一")
    end = model.index("# 附录")
    model = model[start:end].strip()
    refs_start = model.index("# 9 参考文献")
    model_body = model[:refs_start].strip()
    references = model[refs_start:].replace("# 9 参考文献", "# 参考文献", 1).strip()
    return f"{prelim}\n\n{model_body}\n\n{CONCLUSION.strip()}\n\n{references}\n"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--font-size", type=float, default=12.0)
    parser.add_argument("--line-spacing", type=float, default=1.25)
    args = parser.parse_args()

    title, abstract, keywords = parse_front_matter(TITLE_SOURCE.read_text(encoding="utf-8"))
    builder = PaperBuilder(font_size=args.font_size, line_spacing=args.line_spacing)
    builder.add_front_matter(title, abstract, keywords)
    builder.add_markdown(prepare_body())
    builder.save(args.output)
    print(f"Saved: {args.output}")
    print(f"Tables: {builder.table_count}; figures: {builder.figure_count}; equations: {builder.equation_count}")


if __name__ == "__main__":
    main()
