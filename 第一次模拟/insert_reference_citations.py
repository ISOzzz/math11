from pathlib import Path
import re

from docx import Document


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "运行结果_20260819" / "数学建模论文初稿_约25页.docx"
CANDIDATE = ROOT / "运行结果_20260819" / "数学建模论文终稿_引用标注候选.docx"


def find_paragraph(document: Document, needle: str):
    matches = [paragraph for paragraph in document.paragraphs if needle in paragraph.text]
    if len(matches) != 1:
        raise RuntimeError(f"定位段落失败：{needle!r}，匹配数={len(matches)}")
    return matches[0]


def replace_once(paragraph, old: str, new: str) -> None:
    text = paragraph.text
    if text.count(old) != 1:
        raise RuntimeError(f"目标文本匹配数不是1：{old!r}，实际={text.count(old)}")
    paragraph.text = text.replace(old, new, 1)


def main() -> None:
    document = Document(SOURCE)

    literature = find_paragraph(document, "Rijal等指出")
    replace_once(
        literature,
        "不能把工人完全等同于同质化设备[2]。因此",
        "不能把工人完全等同于同质化设备[2]。"
        "Winkelhaus等从Order Picking 4.0视角强调数字技术、作业系统与人员因素的协同演进[3]。因此",
    )

    problem_one = find_paragraph(document, "问题一需要每天选择5个8小时班次")
    replace_once(
        problem_one,
        "使人员分布更加均衡。",
        "使人员分布更加均衡[4]。",
    )

    low_efficiency = find_paragraph(document, "假设4：第二、三问中每名工人的1个低效小时")
    replace_once(
        low_efficiency,
        "把疲劳或休息影响表示为可调度的离散能力损失，使模型仍保持线性",
        "把疲劳或休息影响表示为可调度的离散能力损失[5]，使模型仍保持线性",
    )

    system_assumption = find_paragraph(document, "假设5：货物在当日内可以无损暂存")
    replace_once(
        system_assumption,
        "其影响是所得方案属于理想运行条件下的人员下界",
        "已有研究表明，人机协同拣选系统还可能改变仓库员工的工作体验与负荷[6]。"
        "其影响是所得方案属于理想运行条件下的人员下界",
    )

    modelling = find_paragraph(document, "问题一要求依据逐小时到货量安排每天 5 个班次")
    replace_once(
        modelling,
        "同时尽量减少用工。该问题同时包含",
        "同时尽量减少用工。仓储研究的长期综述表明，运筹优化是仓储系统分析与决策的重要方法[7]。"
        "该问题同时包含",
    )
    replace_once(
        modelling,
        "故建立库存平衡混合整数线性规划模型。",
        "故建立库存平衡混合整数线性规划模型[8]。",
    )

    document.save(CANDIDATE)

    check = Document(CANDIDATE)
    reference_heading = next(
        i for i, paragraph in enumerate(check.paragraphs)
        if paragraph.text.strip() == "参考文献"
    )
    body = "\n".join(paragraph.text for paragraph in check.paragraphs[:reference_heading])
    reference_text = "\n".join(paragraph.text for paragraph in check.paragraphs[reference_heading + 1:])

    missing_body = [number for number in range(1, 9) if f"[{number}]" not in body]
    missing_entries = [number for number in range(1, 9) if f"[{number}]" not in reference_text]
    if missing_body or missing_entries:
        raise RuntimeError(
            f"引用完整性校验失败：正文缺失={missing_body}，参考文献表缺失={missing_entries}"
        )

    occurrences = []
    for index, paragraph in enumerate(check.paragraphs[:reference_heading]):
        labels = re.findall(r"\[([1-8])\]", paragraph.text)
        if labels:
            occurrences.append((index, labels, paragraph.text))

    first_seen = []
    for _, labels, _ in occurrences:
        for label in labels:
            number = int(label)
            if number not in first_seen:
                first_seen.append(number)
    if first_seen != list(range(1, 9)):
        raise RuntimeError(f"首次引用顺序不符合顺序编码制：{first_seen}")

    print(f"CANDIDATE={CANDIDATE}")
    print(f"REFERENCE_HEADING_INDEX={reference_heading}")
    print(f"FIRST_SEEN_ORDER={first_seen}")
    print("CITATION_MAPPING")
    for index, labels, text in occurrences:
        print(f"P{index}: {','.join(labels)} | {text}")


if __name__ == "__main__":
    main()
