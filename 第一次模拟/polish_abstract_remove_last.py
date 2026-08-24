from pathlib import Path

from docx import Document


SOURCE = Path(r"D:\第一次模拟\运行结果_20260819\数学建模论文终稿.docx")
CANDIDATE = Path(r"D:\第一次模拟\运行结果_20260819\数学建模论文终稿_摘要删减候选.docx")

ABSTRACT_PARAGRAPHS = [
    (
        "随着电子商务、即时零售与多渠道履约持续发展，物流分拣中心的小时到货波动、服务时限约束与人工成本矛盾日益突出。"
        "传统经验排班难以同时协调货物清零、效率衰减和连续工日要求，容易造成阶段性积压或人员冗余。"
        "基于未来30天逐小时到货量，本文研究从日内货物处理、低效时段调度到月度人员配置的递进优化问题。"
    ),
    (
        "针对问题一的日内班次与人数优化，建立库存平衡班次优化模型。"
        "以班次启用、班次人数、小时处理量和剩余货量为决策变量，通过库存递推保证货物不被提前处理且日末清零，"
        "并采用词典序目标依次最小化总人日、在岗峰值、小时人员偏差和累计库存。"
        "混合整数线性规划求解表明，全月需10791人日，日均359.70人，处理能力利用率达99.88%；"
        "30天均达到每人每日处理200件对应的理论下界。最大约束残差、整数性误差和最优性间隙均为0，"
        "逐日独立重求解一致率为100%。"
    ),
    (
        "针对问题二的时限与低效约束排班，建立时限与低效约束排班模型。"
        "将货物划分为早期货物和普通货物，引入低效小时人数变量，显式刻画每名工人1小时处理效率由25件/小时降至10件/小时产生的能力损失，"
        "并分别设置两类库存平衡和截止时间约束。求解结果为全月11858人日，较问题一增加9.89%，处理能力利用率为98.26%；"
        "0:00至12:00到达的早期货物均在16:00前清零，全部货物均在24:00前完成，逐日独立重求解一致率为100%。"
    ),
    (
        "针对问题三的月度人员配置，建立连续工日约束人员配置优化模型。"
        "利用“累计工日—连续工日”状态流网络刻画员工出勤状态，在满足每人工作23天且连续工作不超过7天的条件下最小化招聘人数，"
        "并将聚合整数流分解为具名员工月历。结果表明，最少招聘581人，共安排13363人日，全月小时在岗峰值为295人；"
        "理论下界与可行解同为581人，最优性间隙为0。对720小时到货量进行20次独立±10%随机扰动后，"
        "三问目标值最大绝对变化率分别为0.473%、1.240%和1.721%，说明模型对零均值小时噪声具有较强鲁棒性。"
    ),
]


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def main() -> None:
    total_chars = sum(len(text) for text in ABSTRACT_PARAGRAPHS)
    if not 800 <= total_chars <= 1000:
        raise ValueError(f"摘要字符数不符合要求：{total_chars}")

    document = Document(SOURCE)
    heading_index = next(i for i, p in enumerate(document.paragraphs) if p.text.strip() == "摘要")
    keywords_index = next(i for i, p in enumerate(document.paragraphs) if p.text.startswith("关键词："))
    current = document.paragraphs[heading_index + 1 : keywords_index]
    if len(current) < 4:
        raise RuntimeError(f"现有摘要段落不足：{len(current)}")

    for paragraph, text in zip(current[:4], ABSTRACT_PARAGRAPHS):
        paragraph.clear()
        paragraph.add_run(text)

    for paragraph in current[4:]:
        remove_paragraph(paragraph)

    document.save(CANDIDATE)

    check = Document(CANDIDATE)
    heading_index = next(i for i, p in enumerate(check.paragraphs) if p.text.strip() == "摘要")
    keywords_index = next(i for i, p in enumerate(check.paragraphs) if p.text.startswith("关键词："))
    saved = [p.text for p in check.paragraphs[heading_index + 1 : keywords_index]]
    if saved != ABSTRACT_PARAGRAPHS:
        raise RuntimeError("摘要写入或末段删除校验失败")
    if any("本文创新性构建" in text for text in saved):
        raise RuntimeError("独立创新总结段仍然存在")

    print(f"CANDIDATE={CANDIDATE}")
    print(f"ABSTRACT_PARAGRAPHS={len(saved)}")
    print(f"ABSTRACT_CHARS={sum(len(text) for text in saved)}")
    print(f"PARAGRAPH_CHARS={[len(text) for text in saved]}")


if __name__ == "__main__":
    main()
