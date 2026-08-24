from pathlib import Path
import re
import zipfile

from docx import Document


PATH = Path(r"D:\第一次模拟\运行结果_20260819\数学建模论文终稿.docx")


document = Document(PATH)
with zipfile.ZipFile(PATH) as package:
    zip_error = package.testzip()

reference_heading = next(
    i for i, paragraph in enumerate(document.paragraphs)
    if paragraph.text.strip() == "参考文献"
)
body_paragraphs = document.paragraphs[:reference_heading]
reference_paragraphs = document.paragraphs[reference_heading + 1:]

occurrences = []
for index, paragraph in enumerate(body_paragraphs):
    labels = [int(value) for value in re.findall(r"\[([1-8])\]", paragraph.text)]
    if labels:
        occurrences.append((index, labels, paragraph.text))

first_seen = []
for _, labels, _ in occurrences:
    for label in labels:
        if label not in first_seen:
            first_seen.append(label)

reference_entries = {
    int(match.group(1)): paragraph.text
    for paragraph in reference_paragraphs
    if (match := re.match(r"^\[([1-8])\]", paragraph.text.strip()))
}

body_counts = {
    number: sum(labels.count(number) for _, labels, _ in occurrences)
    for number in range(1, 9)
}

print(f"ZIP_TEST={zip_error}")
print(f"PARAGRAPHS={len(document.paragraphs)}")
print(f"TABLES={len(document.tables)}")
print(f"REFERENCE_ENTRIES={sorted(reference_entries)}")
print(f"FIRST_SEEN_ORDER={first_seen}")
print(f"BODY_CITATION_COUNTS={body_counts}")
print(f"CITATION_PARAGRAPHS={[index for index, _, _ in occurrences]}")

assert zip_error is None
assert sorted(reference_entries) == list(range(1, 9))
assert first_seen == list(range(1, 9))
assert all(count >= 1 for count in body_counts.values())
assert len(document.tables) == 8
