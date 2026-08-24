from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"D:\第一次模拟\运行结果_20260819\数学建模论文终稿_摘要二次优化.docx")
OUTPUT = Path(r"D:\第一次模拟\运行结果_20260819\数学建模论文终稿_模块化摘要.docx")

ABSTRACT_PARAGRAPHS = [
    (
        "随着电子商务与即时配送规模持续扩张，物流分拣中心面临到货量小时波动、处理时限趋严与用工成本上升等挑战。"
        "传统经验排班难以同时协调库存清零、效率衰减和连续工日限制，易造成货物积压或人员冗余。"
        "为此，本文基于未来30天逐小时到货量，研究日内货物处理、低效时段安排与月度人员配置的递进优化问题。"
    ),
    (
        "针对问题一的日内班次与人数优化，建立库存平衡班次优化模型，以班次启用、班次人数、小时处理量和库存量为决策变量，"
        "通过库存递推约束禁止提前处理，并采用词典序目标依次最小化总人日、在岗峰值、小时人员偏差和累计积压。"
        "混合整数线性规划求解结果表明，全月需10791人日，日均359.70人，处理能力利用率为99.88%；"
        "30天均达到每人每日处理200件对应的理论下界，最大约束残差和整数性误差均为0，逐日独立重求解一致率为100%。"
    ),
    (
        "针对问题二的时限与低效约束排班，将货物划分为早期货物和普通货物，引入低效小时人数变量，"
        "显式刻画每名工人1小时效率降至10件/小时产生的能力损失，并分别设置库存平衡与截止时间约束。"
        "求解得到全月11858人日，较问题一增加9.89%，处理能力利用率为98.26%；"
        "0:00（含）至12:00（不含）的早期货物均在16:00前清零，全部货物均在24:00前完成，30天独立重求解一致率为100%。"
    ),
    (
        "针对问题三的月度人员配置，建立连续工日约束人员配置优化模型，以“累计工日—连续工日”状态流网络描述员工出勤状态，"
        "在满足每人工作23天、连续工作不超过7天的条件下最小化招聘人数，并将聚合流分解为具名员工月历。"
        "结果表明，最少招聘581人，共安排13363人日，全月小时在岗峰值为295人；理论下界与可行解均为581人，最优性间隙为0。"
        "对720小时到货量进行20次独立±10%随机扰动后，三问目标值最大绝对变化率依次为0.473%、1.240%和1.721%，"
        "验证了模型对零均值小时噪声的鲁棒性。"
    ),
    (
        "本文创新性构建“班次—低效小时—员工月历”多层递进优化框架，显式刻画低效能力损失，"
        "并以状态流分解实现聚合需求向个体排班的转化，可为物流分拣中心的班次设计、人员招聘与轮休管理提供定量决策依据。"
    ),
]


def format_abstract_paragraph(paragraph, text: str, *, keep_with_next: bool) -> None:
    paragraph.style = "Normal"
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "宋体")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")

    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.widow_control = True


def main() -> None:
    total_chars = sum(len(text) for text in ABSTRACT_PARAGRAPHS)
    if not 800 <= total_chars <= 1000:
        raise ValueError(f"摘要字符数不符合要求：{total_chars}")
    if not all(
        ABSTRACT_PARAGRAPHS[i].startswith(f"针对问题{'一二三'[i - 1]}")
        for i in range(1, 4)
    ):
        raise ValueError("问题段落未按指定格式起始")

    document = Document(SOURCE)
    if len(document.paragraphs) < 4 or document.paragraphs[1].text.strip() != "摘要":
        raise RuntimeError("未能定位终稿摘要标题")

    keywords = next(
        (paragraph for paragraph in document.paragraphs if paragraph.text.startswith("关键词：")),
        None,
    )
    if keywords is None:
        raise RuntimeError("未能定位关键词段落")

    format_abstract_paragraph(document.paragraphs[2], ABSTRACT_PARAGRAPHS[0], keep_with_next=False)
    for index, text in enumerate(ABSTRACT_PARAGRAPHS[1:], start=1):
        paragraph = keywords.insert_paragraph_before(style="Normal")
        format_abstract_paragraph(
            paragraph,
            text,
            keep_with_next=index == len(ABSTRACT_PARAGRAPHS) - 1,
        )

    document.save(OUTPUT)

    check = Document(OUTPUT)
    abstract_heading_index = next(i for i, p in enumerate(check.paragraphs) if p.text.strip() == "摘要")
    keywords_index = next(i for i, p in enumerate(check.paragraphs) if p.text.startswith("关键词："))
    saved_paragraphs = [p.text for p in check.paragraphs[abstract_heading_index + 1 : keywords_index]]
    if saved_paragraphs != ABSTRACT_PARAGRAPHS:
        raise RuntimeError("摘要分段写入后校验失败")

    print(f"OUTPUT={OUTPUT}")
    print(f"ABSTRACT_PARAGRAPHS={len(saved_paragraphs)}")
    print(f"ABSTRACT_CHARS={sum(len(text) for text in saved_paragraphs)}")
    print(f"PARAGRAPH_CHARS={[len(text) for text in saved_paragraphs]}")


if __name__ == "__main__":
    main()
